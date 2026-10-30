"""End-to-end smoke tests -- actually invoke each package's CLI as a
subprocess (`python3 -m pt.cli ...`, same as a real user would), against
test_data/'s synthetic household, and just check the exit code and a
couple of expected output fragments.

This layer exists because tests/test_import.py, tests/test_projection.py,
and tests/test_roth_optimizer.py all call internal functions directly --
precise and fast, but blind to a broken argparse wiring (a wrong flag
name, a positional/optional ordering bug) since they never go through
cli.py's actual argument parsing at all. The --account flag's "must come
before the file list" behavior (an argparse quirk with the existing
nargs="+" positional -- see pt/cli.py's import subcommand) is exactly the
kind of thing only a real CLI invocation catches; that's the gap this
file covers. Keep it to a handful of "does the command still basically
work" checks, not a copy of the detailed internal-function tests above.
"""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from . import _helpers as h

REPO_ROOT = h.REPO_ROOT


def run_cli(module: str, *args, cwd=REPO_ROOT, timeout=60):
    """Runs `python3 -m {module} {args}` as a real subprocess (the exact
    way a user invokes any of this project's tools), from the repo root
    (so relative test_data/ paths resolve the same way the examples in
    README.md/test_data/README.md do). Returns the completed process."""
    return subprocess.run(
        [sys.executable, "-m", module, *args],
        cwd=cwd, capture_output=True, text=True, timeout=timeout,
    )


class PtCliSmokeTests(unittest.TestCase):
    def setUp(self):
        self._db = h.TempDB()
        self.db_path = self._db.__enter__()
        conn = h.seeded_connection(self.db_path)
        h.import_fidelity_household(conn)
        conn.close()

    def tearDown(self):
        self._db.__exit__(None, None, None)

    def test_project_runs_cleanly(self):
        result = run_cli(
            "pt.cli", "--db", str(self.db_path), "project",
            "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE),
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Blended growth rate", result.stdout)

    def test_monte_carlo_runs_cleanly(self):
        result = run_cli(
            "pt.cli", "--db", str(self.db_path), "monte-carlo",
            "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE),
            "--trials", "20", "--seed", "1",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Success rate", result.stdout)

    def test_project_per_bucket_growth_prints_a_breakdown(self):
        result = run_cli(
            "pt.cli", "--db", str(self.db_path), "project",
            "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE),
            "--per-bucket-growth",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("household summary", result.stdout)
        self.assertIn("Per-bucket growth rate", result.stdout)
        self.assertIn("Taxable", result.stdout)

    def test_monte_carlo_per_bucket_growth_runs_cleanly(self):
        result = run_cli(
            "pt.cli", "--db", str(self.db_path), "monte-carlo",
            "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE),
            "--trials", "20", "--seed", "1", "--per-bucket-growth",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Success rate", result.stdout)
        self.assertIn("Per-bucket mean blended return", result.stdout)

    def test_report_runs_cleanly_and_embeds_projection_tab(self):
        with tempfile.TemporaryDirectory() as out_dir:
            xlsx = Path(out_dir) / "r.xlsx"
            txt = Path(out_dir) / "r.txt"
            result = run_cli("pt.cli", "--db", str(self.db_path), "report",
                              "--out", str(xlsx), "--text-out", str(txt),
                              "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE))
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(xlsx.exists())
            self.assertTrue(txt.exists())
            # The workbook's Projection tab is only added when the scenario
            # projects cleanly -- and report.py must feed compute_blended_
            # growth_rate() the subclass-level allocation rows, not the
            # macro-level ones, or a scenario using the newer
            # returns.asset_classes shape (test_data/'s does) silently
            # drops the tab.
            import openpyxl
            wb = openpyxl.load_workbook(xlsx)
            self.assertIn("Projection", wb.sheetnames)

    def test_report_monte_carlo_flag_adds_tab_right_after_projection(self):
        with tempfile.TemporaryDirectory() as out_dir:
            xlsx = Path(out_dir) / "r.xlsx"
            txt = Path(out_dir) / "r.txt"
            result = run_cli("pt.cli", "--db", str(self.db_path), "report",
                              "--out", str(xlsx), "--text-out", str(txt),
                              "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE),
                              "--monte-carlo", "--trials", "20", "--seed", "1")
            self.assertEqual(result.returncode, 0, result.stderr)
            import openpyxl
            names = openpyxl.load_workbook(xlsx).sheetnames
            self.assertIn("Monte Carlo", names)
            self.assertEqual(names.index("Monte Carlo"), names.index("Projection") + 1)

    def test_report_per_bucket_growth_adds_a_breakdown_to_both_tabs(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as out_dir:
            xlsx = Path(out_dir) / "r.xlsx"
            txt = Path(out_dir) / "r.txt"
            result = run_cli("pt.cli", "--db", str(self.db_path), "report", "--private",
                              "--out", str(xlsx), "--text-out", str(txt),
                              "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE),
                              "--monte-carlo", "--trials", "20", "--seed", "1", "--per-bucket-growth")
            self.assertEqual(result.returncode, 0, result.stderr)
            wb = openpyxl.load_workbook(xlsx)
            for sheet_name in ("Projection", "Monte Carlo"):
                col_a = [r[0] for r in wb[sheet_name].iter_rows(min_col=1, max_col=1, values_only=True)]
                self.assertTrue(any(v and "Per-bucket growth rate" in str(v) for v in col_a), sheet_name)
                self.assertTrue(any(v and str(v).strip().startswith("Taxable:") for v in col_a), sheet_name)

    HYPOTHETICAL_SCENARIO = """
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns:
  asset_classes:
    US Equity: {expected_return: 7.5, volatility: 18}
    US Bonds: {expected_return: 4.0, volatility: 6}
hypothetical_accounts:
  - account_number: taxable1
    account_type: Brokerage
    owner: alex
    value: 500000
    asset_classes: {US Equity: 100}
  - account_number: ira1
    account_type: IRA
    owner: alex
    value: 250000
    asset_classes: {US Equity: 40, US Bonds: 60}
"""

    def test_report_hypothetical_scenario_uses_its_own_accounts_not_the_snapshot(self):
        """If --scenario configures hypothetical_accounts, the WHOLE report
        is about that made-up household, not self.db_path's real (and
        completely unrelated) test-fidelity snapshot -- see
        build_workbook()'s docstring. The Projection tab's growth math
        (including --per-bucket-growth) comes from it -- a 100%-US-Equity
        taxable account blends to exactly that class's own expected_return
        (7.5%) -- and so do the Summary/Accounts/Allocation tabs and the
        text report, while the holdings-level tabs are left out."""
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "hypothetical.yaml"
            scenario_path.write_text(self.HYPOTHETICAL_SCENARIO)
            xlsx = Path(d) / "r.xlsx"
            txt = Path(d) / "r.txt"
            result = run_cli("pt.cli", "--db", str(self.db_path), "report", "--private",
                              "--out", str(xlsx), "--text-out", str(txt),
                              "--scenario", str(scenario_path), "--per-bucket-growth")
            self.assertEqual(result.returncode, 0, result.stderr)
            wb = openpyxl.load_workbook(xlsx)
            col_a = [v for v in wb["Projection"].iter_rows(min_col=1, max_col=1, values_only=True) for v in v]
            self.assertTrue(any(v and str(v).strip() == "Taxable: 7.50%" for v in col_a), col_a)

            # No holdings exist for a hypothetical account.
            for dropped in ("Holdings", "Holdings by Ticker", "Investment Expenses", "Unclassified"):
                self.assertNotIn(dropped, wb.sheetnames)
            self.assertIn("Hypothetical household from scenario: hypothetical.yaml", wb["Summary"]["A2"].value)
            self.assertIn("$750,000", wb["Summary"]["A3"].value)

            # Accounts: the hypothetical ids, with no separate Account Number column.
            accounts = list(wb["Accounts"].iter_rows(values_only=True))
            self.assertEqual(list(accounts[0][:4]), ["Account Name", "Owner", "Account Type", "Value"])
            self.assertEqual([(r[0], r[3]) for r in accounts[1:]], [("taxable1", 500000), ("ira1", 250000)])

            # Allocation: each account's own asset-class mix.
            alloc_rows = [r for r in wb["Allocation"].iter_rows(min_row=4, values_only=True) if r[0]]
            self.assertEqual({r[1] for r in alloc_rows}, {"alex"})
            ira = {r[4]: (r[3], r[5], r[6]) for r in alloc_rows if r[0] == "ira1"}
            self.assertEqual(ira["US Bonds"], ("Income", 150000, 0.6))
            self.assertEqual(ira["US Equity"], ("Growth", 100000, 0.4))

            body = txt.read_text()
            self.assertIn("Hypothetical household from scenario: hypothetical.yaml", body)
            self.assertNotIn("test-fidelity", body)
            self.assertNotIn("INVESTMENT EXPENSE SUMMARY", body)

    def test_report_hypothetical_scenario_needs_no_snapshot(self):
        """A hypothetical report never touches the database, so it works
        before anything has been imported."""
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "hypothetical.yaml"
            scenario_path.write_text(self.HYPOTHETICAL_SCENARIO)
            xlsx = Path(d) / "r.xlsx"
            result = run_cli("pt.cli", "--db", str(Path(d) / "empty.db"), "report", "--private",
                              "--out", str(xlsx), "--text-out", str(Path(d) / "r.txt"),
                              "--scenario", str(scenario_path))
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Projection", openpyxl.load_workbook(xlsx).sheetnames)

    def test_report_hypothetical_scenario_with_a_bad_account_fails_clearly(self):
        """A broken hypothetical_accounts section is an error, not a silent
        fallback to the real accounts."""
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "hypothetical.yaml"
            scenario_path.write_text(self.HYPOTHETICAL_SCENARIO.replace("account_type: IRA", "account_type: Nope"))
            result = run_cli("pt.cli", "--db", str(self.db_path), "report", "--private",
                              "--out", str(Path(d) / "r.xlsx"), "--text-out", str(Path(d) / "r.txt"),
                              "--scenario", str(scenario_path))
            self.assertEqual(result.returncode, 1)
            self.assertIn("hypothetical_accounts", result.stderr)

    def test_report_public_is_the_default_and_trims_the_shareable_view(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as out_dir:
            xlsx = Path(out_dir) / "r.xlsx"
            txt = Path(out_dir) / "r.txt"
            result = run_cli("pt.cli", "--db", str(self.db_path), "report",
                              "--out", str(xlsx), "--text-out", str(txt),
                              "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE))
            self.assertEqual(result.returncode, 0, result.stderr)
            names = openpyxl.load_workbook(xlsx).sheetnames
            for dropped in ("Holdings", "Holdings by Ticker", "Investment Expenses", "IPS Comparison",
                            "IPS Comparison (Macro)", "IPS by Account", "Unclassified"):
                self.assertNotIn(dropped, names)
            # Account Number column gone; decedent dates present (test_data's
            # scenario has inherited_accounts tied to held accounts).
            acct_header = [c.value for c in openpyxl.load_workbook(xlsx)["Accounts"][1]]
            self.assertNotIn("Account Number", acct_header)
            self.assertIn("Decedent Date of Death", acct_header)
            body = txt.read_text()
            self.assertNotIn("#...", body)
            self.assertNotIn("IPS TARGET", body)
            self.assertNotIn("EXPENSE SUMMARY", body)
            # source_file ("test-fidelity", see _helpers.import_fidelity_household)
            # is redacted to a file count, not shown verbatim.
            self.assertNotIn("test-fidelity", body)
            self.assertIn("<redacted> (1 file)", body)
            summary_a2 = openpyxl.load_workbook(xlsx)["Summary"]["A2"].value
            self.assertNotIn("test-fidelity", summary_a2)
            self.assertIn("<redacted>", summary_a2)

    def test_report_private_keeps_every_tab_and_the_account_number_column(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as out_dir:
            xlsx = Path(out_dir) / "r.xlsx"
            txt = Path(out_dir) / "r.txt"
            result = run_cli("pt.cli", "--db", str(self.db_path), "report", "--private",
                              "--out", str(xlsx), "--text-out", str(txt),
                              "--scenario", str(h.TEST_SCENARIO), "--profile", str(h.TEST_PROFILE))
            self.assertEqual(result.returncode, 0, result.stderr)
            wb = openpyxl.load_workbook(xlsx)
            for kept in ("Holdings", "Investment Expenses", "IPS Comparison", "IPS Comparison (Macro)", "Unclassified"):
                self.assertIn(kept, wb.sheetnames)
            # Holdings by Ticker sits right after Holdings
            self.assertEqual(wb.sheetnames.index("Holdings by Ticker"), wb.sheetnames.index("Holdings") + 1)

            # Holdings: % of Portfolio sums to 100% overall, % of Account to
            # 100% within each account. Holdings has no account-number column
            # and two test accounts share a name ("ROLLOVER IRA"), so each
            # name's shares sum to 100% x the number of accounts with it.
            hold = list(wb["Holdings"].iter_rows(values_only=True))
            hdr = list(hold[0])
            self.assertEqual(hdr[hdr.index("Current Value") + 1:hdr.index("Current Value") + 3],
                             ["% of Portfolio", "% of Account"])
            self.assertEqual(hdr[hdr.index("% of Account") + 1], "Asset Class")
            rows = hold[1:]
            port_col, acct_col, acct_name_col = hdr.index("% of Portfolio"), hdr.index("% of Account"), 0
            self.assertAlmostEqual(sum(r[port_col] for r in rows), 1.0, places=6)
            by_acct = {}
            for r in rows:
                by_acct[r[acct_name_col]] = by_acct.get(r[acct_name_col], 0.0) + r[acct_col]
            acct_tab = list(wb["Accounts"].iter_rows(values_only=True))
            name_col = acct_tab[0].index("Account Name")
            accounts_named = {}
            for r in acct_tab[1:]:
                accounts_named[r[name_col]] = accounts_named.get(r[name_col], 0) + 1
            for name, share in by_acct.items():
                self.assertAlmostEqual(share, accounts_named[name], places=6, msg=name)

            # Holdings by Ticker: one row per symbol whose values add up to
            # the per-position values on Holdings, then a TOTAL row.
            tick = list(wb["Holdings by Ticker"].iter_rows(values_only=True))
            self.assertEqual(list(tick[0][:5]), ["Symbol", "Description", "Total Quantity", "Total Value",
                                                 "% of Portfolio"])
            tick_rows, total = tick[1:-1], tick[-1]
            self.assertEqual(total[0], "TOTAL")
            sym_col, val_col = hdr.index("Symbol"), hdr.index("Current Value")
            expected = {}
            for r in rows:
                expected[r[sym_col]] = expected.get(r[sym_col], 0.0) + r[val_col]
            self.assertEqual({r[0] for r in tick_rows}, set(expected))
            for r in tick_rows:
                self.assertAlmostEqual(r[3], expected[r[0]], places=2)
            self.assertEqual([r[3] for r in tick_rows], sorted((r[3] for r in tick_rows), reverse=True))
            self.assertAlmostEqual(sum(r[4] for r in tick_rows), 1.0, places=6)
            self.assertAlmostEqual(total[3], sum(expected.values()), places=2)

            acct_header = [c.value for c in wb["Accounts"][1]]
            self.assertEqual(acct_header[0], "Account Number")
            self.assertIn("Decedent Date of Death", acct_header)
            # account numbers show only their last 4 digits in the private view
            acct_nums = [r[0] for r in wb["Accounts"].iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
            self.assertTrue(acct_nums)
            for num in acct_nums:
                self.assertTrue(len(num) <= 4 or (num.startswith("...") and len(num) == 7), num)
            # ...and the scenario dump masks them the same way, not "<redacted>"
            proj_col_a = [r[0] for r in wb["Projection"].iter_rows(min_col=1, max_col=1, values_only=True)]
            acct_lines = [str(v) for v in proj_col_a if v and "account_number:" in str(v)]
            self.assertTrue(acct_lines)
            self.assertFalse(any("<redacted>" in v for v in acct_lines))
            for line in acct_lines:
                value = line.split("account_number:", 1)[1].strip().strip("'\"")
                self.assertTrue(len(value) <= 4 or value.startswith("..."), line)
            self.assertIn("test-fidelity", wb["Summary"]["A2"].value)
            self.assertIn("test-fidelity", txt.read_text())

    def test_import_with_missing_etrade_account_fails_clearly(self):
        """The exact ordering/flag issue caught by hand this session --
        --account must be accepted and, when omitted for an E*TRADE file,
        must fail with a clear message rather than a traceback."""
        result = run_cli("pt.cli", "--db", str(self.db_path), "import", str(h.ETRADE_CASEY))
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("account number", result.stderr.lower())

    def test_import_multi_broker_with_account_flag_succeeds(self):
        result = run_cli(
            "pt.cli", "--db", str(self.db_path), "import",
            "--account", f"{h.ETRADE_CASEY.name}:{h.ETRADE_CASEY_ACCOUNT}",
            str(h.SCHWAB_JORDAN), "Jordan", str(h.ETRADE_CASEY), "Casey",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Imported snapshot", result.stdout)


class RothOptimizerCliSmokeTests(unittest.TestCase):
    def test_optimize_runs_cleanly_against_hypothetical_scenario(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 500000
    pct_growth: 60
roth_conversions:
  - owner: alex
    start_date: 2027-01-01
    end_date: 2033-12-31
    constraints: {max_marginal_bracket: 24, max_conversion_per_year: 60000}
""")
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "2",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Optimized schedule", result.stdout)

    _SCATTER_SCENARIO = """
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 500000
    pct_growth: 60
roth_conversions:
  - owner: alex
    start_date: 2027-01-01
    end_date: 2033-12-31
    constraints: {max_marginal_bracket: 24, max_conversion_per_year: 60000}
"""

    def test_optimize_records_the_html_chart_on_the_saved_scenarios_roth_conversions(self):
        import yaml
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text(self._SCATTER_SCENARIO)
            out_scenario = Path(d) / "out" / "scenario_optimized_24.yaml"
            chart = Path(d) / "out" / "charts" / "search.html"
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "2",
                "--out-scenario", str(out_scenario),
                "--save-iterations", str(Path(d) / "out" / "search.csv"),
                "--save-iterations", str(chart),
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            saved = yaml.safe_load(out_scenario.read_text())
            # Relative to the saved scenario file's own directory, so the
            # two can be moved together; only the .html counts (not the .csv).
            self.assertEqual(saved["roth_conversions"][0]["scatter_chart"], "charts/search.html")
            self.assertTrue(chart.exists())

    def test_optimize_without_an_html_chart_records_no_scatter_chart_field(self):
        import yaml
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text(self._SCATTER_SCENARIO)
            out_scenario = Path(d) / "scenario_optimized_24.yaml"
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "2",
                "--out-scenario", str(out_scenario),
                "--save-iterations", str(Path(d) / "search.csv"),
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            saved = yaml.safe_load(out_scenario.read_text())
            self.assertNotIn("scatter_chart", saved["roth_conversions"][0])

    def test_optimize_per_bucket_growth_prints_a_breakdown(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns:
  asset_classes:
    US Equity: {expected_return: 8.0, volatility: 18}
    US Bonds: {expected_return: 3.0, volatility: 7}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 500000
    asset_classes: {US Bonds: 100}
  - account_number: roth1
    account_type: Roth IRA
    owner: alex
    value: 50000
    asset_classes: {US Equity: 100}
roth_conversions:
  - owner: alex
    start_date: 2027-01-01
    end_date: 2033-12-31
    constraints: {max_marginal_bracket: 24, max_conversion_per_year: 60000}
""")
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "2", "--per-bucket-growth",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Optimized schedule", result.stdout)
            self.assertIn("Per-bucket growth rates", result.stdout)
            self.assertIn("Traditional (Alex: IRA): 3.00%", result.stdout)
            self.assertIn("Roth (Alex): 8.00%", result.stdout)

    def test_sweep_prints_a_comparison_table_without_declaring_a_winner(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 500000
    pct_growth: 60
roth_conversions:
  - owner: alex
    start_date: 2027-01-01
    end_date: 2033-12-31
    constraints: {max_marginal_bracket: 24, max_conversion_per_year: 60000}
""")
            result = run_cli(
                "roth_optimizer.cli", "sweep", "--scenario", str(scenario_path),
                "--heir-tax-rates", "12,24,35", "--seed", "1", "--restarts", "1",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Heir tax rate sweep", result.stdout)
            for rate in ("12%", "24%", "35%"):
                self.assertIn(rate, result.stdout)
            # Deliberately no single "winner" declared -- see cmd_sweep()'s
            # docstring for why (--heir-tax-rate is an assumption, not
            # something the search optimizes; the raw dollar totals aren't
            # comparable across different rate assumptions).
            self.assertNotIn("Best rate", result.stdout)
            self.assertNotIn("Optimal rate", result.stdout)

    def test_sweep_out_dir_saves_one_file_pair_per_rate(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "my_scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 500000
    pct_growth: 60
roth_conversions:
  - owner: alex
    start_date: 2027-01-01
    end_date: 2033-12-31
    constraints: {max_marginal_bracket: 24, max_conversion_per_year: 60000}
""")
            out_dir = Path(d) / "out"
            result = run_cli(
                "roth_optimizer.cli", "sweep", "--scenario", str(scenario_path),
                "--heir-tax-rates", "12,24", "--seed", "1", "--restarts", "1", "--out-dir", str(out_dir),
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            for rate_tag in ("12", "24"):
                self.assertTrue((out_dir / f"my_scenario_optimized_{rate_tag}.yaml").exists())
                self.assertTrue((out_dir / f"my_scenario_optimized_{rate_tag}_projection.txt").exists())

    def test_priority_compares_both_orderings_and_saves_both_variants(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "my_scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
  spouse: {name: sam, birthdate: 1962-01-01}
retirement: {retirement_age: 90, spouse_retirement_age: 90}
income: {target_spending: 0}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: alex_trad
    account_type: IRA
    owner: alex
    value: 2000000
    pct_growth: 0
  - account_number: sam_trad
    account_type: IRA
    owner: sam
    value: 2000000
    pct_growth: 0
roth_conversions:
  - owner: alex
    start_date: 2026-01-01
    end_date: 2026-12-31
    annual_amount: 50000
    constraints: {max_marginal_bracket: 24}
  - owner: sam
    start_date: 2026-01-01
    end_date: 2026-12-31
    annual_amount: 2000000
    constraints: {max_marginal_bracket: 24}
""")
            out_dir = Path(d) / "out"
            result = run_cli(
                "roth_optimizer.cli", "priority", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "1", "--out-dir", str(out_dir),
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Conversion priority comparison", result.stdout)
            self.assertIn("alex prioritized:", result.stdout)
            self.assertIn("sam prioritized:", result.stdout)
            for owner in ("alex", "sam"):
                self.assertTrue((out_dir / f"my_scenario_priority_{owner}first.yaml").exists())
                self.assertTrue((out_dir / f"my_scenario_priority_{owner}first_projection.txt").exists())

    def test_priority_requires_exactly_two_owners(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 500000
    pct_growth: 60
roth_conversions:
  - owner: alex
    start_date: 2027-01-01
    end_date: 2033-12-31
    constraints: {max_marginal_bracket: 24}
""")
            result = run_cli(
                "roth_optimizer.cli", "priority", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24",
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("exactly two", result.stderr)

    def test_optimize_ira_distributions_jointly_searches_both_kinds(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 90}
income: {target_spending: 0}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 1000000
    pct_growth: 0
roth_conversions:
  - owner: alex
    start_date: 2026-01-01
    end_date: 2030-12-31
    constraints: {max_marginal_bracket: 24, max_conversion_per_year: 2000000}
ira_distributions:
  - owner: alex
    start_date: 2026-01-01
    end_date: 2030-12-31
    constraints: {max_marginal_bracket: 24, max_distribution_per_year: 40000}
""")
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "1", "--optimize-ira-distributions",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("alex (Roth conversion,", result.stdout)
            self.assertIn("alex (IRA distribution,", result.stdout)
            self.assertIn("No Roth conversions or IRA distributions at all:", result.stdout)

    def test_optimize_without_the_flag_ignores_ira_distributions(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 90}
income: {target_spending: 0}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 1000000
    pct_growth: 0
roth_conversions:
  - owner: alex
    start_date: 2026-01-01
    end_date: 2030-12-31
    constraints: {max_marginal_bracket: 24, max_conversion_per_year: 2000000}
ira_distributions:
  - owner: alex
    start_date: 2026-01-01
    end_date: 2030-12-31
    constraints: {max_marginal_bracket: 24, max_distribution_per_year: 40000}
""")
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "1",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("alex (Roth conversion,", result.stdout)
            self.assertNotIn("IRA distribution", result.stdout)

    def test_optimize_inherited_withdrawals_searches_the_planned_withdrawal_window(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 90}
income: {target_spending: 0}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: inh1
    account_type: Inherited IRA
    owner: alex
    value: 500000
    pct_growth: 0
inherited_accounts:
  - account_number: inh1
    type: traditional
    decedent_birthdate: 1940-01-01
    decedent_death_date: 2025-06-01
    beneficiary: alex
    planned_withdrawal:
      start_date: 2026-01-01
      end_date: 2034-12-31
      annual_amount: 0
""")
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "1",
                "--optimize-inherited-withdrawals",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("alex (Inherited withdrawal, account inh1,", result.stdout)
            self.assertIn("No Roth conversions or extra inherited-account withdrawals at all:", result.stdout)

    def test_optimize_without_the_flag_ignores_inherited_accounts_planned_withdrawal(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 65}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 500000
    pct_growth: 60
  - account_number: inh1
    account_type: Inherited IRA
    owner: alex
    value: 500000
    pct_growth: 0
roth_conversions:
  - owner: alex
    start_date: 2027-01-01
    end_date: 2033-12-31
    constraints: {max_marginal_bracket: 24}
inherited_accounts:
  - account_number: inh1
    type: traditional
    decedent_birthdate: 1940-01-01
    decedent_death_date: 2025-06-01
    beneficiary: alex
    planned_withdrawal:
      start_date: 2026-01-01
      end_date: 2034-12-31
      annual_amount: 0
""")
            result = run_cli(
                "roth_optimizer.cli", "optimize", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--seed", "1", "--restarts", "1",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("alex (Roth conversion,", result.stdout)
            self.assertNotIn("Inherited withdrawal", result.stdout)

    def test_claim_age_searches_pia_monthly_and_declares_a_winner(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 62}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 800000
    pct_growth: 60
social_security:
  alex: {pia_monthly: 2800}
""")
            result = run_cli(
                "roth_optimizer.cli", "claim-age", "--scenario", str(scenario_path),
                "--heir-tax-rate", "24", "--claim-ages", "62,70",
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Best:", result.stdout)
            self.assertIn("alex", result.stdout)
            self.assertIn("no retirement.life_expectancy", result.stdout)

    def test_claim_age_requires_pia_monthly(self):
        with tempfile.TemporaryDirectory() as d:
            scenario_path = Path(d) / "scenario.yaml"
            scenario_path.write_text("""
people:
  primary: {name: alex, birthdate: 1962-01-01}
retirement: {retirement_age: 62}
income: {target_spending: 60000}
inflation: {general: 2.5}
returns: {growth: 8.0, fixed_income: 4.0}
hypothetical_accounts:
  - account_number: trad1
    account_type: IRA
    owner: alex
    value: 800000
    pct_growth: 60
social_security:
  alex: {monthly_benefit: 2000, claim_age: 67}
""")
            result = run_cli(
                "roth_optimizer.cli", "claim-age", "--scenario", str(scenario_path), "--heir-tax-rate", "24",
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("pia_monthly", result.stderr)


class ScenarioEngineCliSmokeTests(unittest.TestCase):
    def test_generate_and_run_against_test_data_scenario(self):
        with tempfile.TemporaryDirectory() as d:
            sweep_path = Path(d) / "sweep.yaml"
            scenarios_dir = Path(d) / "scenarios"
            results_dir = Path(d) / "results"
            sweep_path.write_text(f"""
base_scenario: {h.TEST_SCENARIO}
parameters:
  retirement.retirement_age: {{min: 63, max: 65, step: 2}}
""")
            gen = run_cli("scenario_engine.cli", "generate", "--sweep", str(sweep_path),
                          "--out-dir", str(scenarios_dir))
            self.assertEqual(gen.returncode, 0, gen.stderr)
            self.assertTrue(list(scenarios_dir.glob("*.yaml")))

            db_path = self._import_household(d)
            run = run_cli("scenario_engine.cli", "run", "--db", str(db_path),
                          "--scenarios-dir", str(scenarios_dir), "--out-dir", str(results_dir),
                          "--profile", str(h.TEST_PROFILE))
            self.assertEqual(run.returncode, 0, run.stderr)
            self.assertTrue((results_dir / "summary.csv").exists())

    def _import_household(self, tmp_dir):
        db_path = Path(tmp_dir) / "scenario_engine_test.db"
        conn = h.seeded_connection(db_path)
        h.import_fidelity_household(conn)
        conn.close()
        return db_path


if __name__ == "__main__":
    unittest.main()
